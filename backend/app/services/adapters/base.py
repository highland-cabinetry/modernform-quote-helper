"""Shared types for document adapters.

An adapter turns one uploaded file into zero or more *parsed orders*. It never
touches the database and never guesses: a row it cannot interpret is reported in
``ignored`` with a reason, and a file it cannot read is reported with a status and
message. The import service decides how parsed orders become stored orders.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from ...models import PackingState

STATE_NAMES = {
    "ALABAMA": "AL", "ALASKA": "AK", "ARIZONA": "AZ", "ARKANSAS": "AR", "CALIFORNIA": "CA", "COLORADO": "CO",
    "CONNECTICUT": "CT", "DELAWARE": "DE", "FLORIDA": "FL", "GEORGIA": "GA", "HAWAII": "HI", "IDAHO": "ID",
    "ILLINOIS": "IL", "INDIANA": "IN", "IOWA": "IA", "KANSAS": "KS", "KENTUCKY": "KY", "LOUISIANA": "LA",
    "MAINE": "ME", "MARYLAND": "MD", "MASSACHUSETTS": "MA", "MICHIGAN": "MI", "MINNESOTA": "MN",
    "MISSISSIPPI": "MS", "MISSOURI": "MO", "MONTANA": "MT", "NEBRASKA": "NE", "NEVADA": "NV",
    "NEW HAMPSHIRE": "NH", "NEW JERSEY": "NJ", "NEW MEXICO": "NM", "NEW YORK": "NY", "NORTH CAROLINA": "NC",
    "NORTH DAKOTA": "ND", "OHIO": "OH", "OKLAHOMA": "OK", "OREGON": "OR", "PENNSYLVANIA": "PA",
    "RHODE ISLAND": "RI", "SOUTH CAROLINA": "SC", "SOUTH DAKOTA": "SD", "TENNESSEE": "TN", "TEXAS": "TX",
    "UTAH": "UT", "VERMONT": "VT", "VIRGINIA": "VA", "WASHINGTON": "WA", "WEST VIRGINIA": "WV",
    "WISCONSIN": "WI", "WYOMING": "WY",
}
_STATE_TOKEN = "|".join(
    sorted([re.escape(name.title()) for name in STATE_NAMES] + [r"[A-Z]{2}"], key=len, reverse=True)
)
CITY_STATE_ZIP = re.compile(
    rf"(?P<city>[A-Za-z][A-Za-z .'-]*?),?\s+(?P<state>{_STATE_TOKEN})[.,]?\s+(?P<zip>\d{{5}})(?:-\d{{4}})?\b",
    re.IGNORECASE,
)


@dataclass(slots=True)
class ParsedLine:
    original_sku: str
    quantity: int
    description: str = ""
    packing_state: PackingState | None = None
    length: float | None = None
    width: float | None = None
    height: float | None = None
    weight: float | None = None
    line_number: int | None = None
    source_ref: str = ""


@dataclass(slots=True)
class ParsedOrder:
    order_number: str | None
    document_type: str = "unknown"
    references: list[str] = field(default_factory=list)
    customer_name: str = ""
    po_number: str = ""
    order_date: str = ""
    ready_date: str = ""
    sales_rep: str = ""
    ship_to: dict[str, str] = field(default_factory=dict)
    residential: bool | None = None
    liftgate: bool | None = None
    packing_state: PackingState = PackingState.FLATPACK
    lines: list[ParsedLine] = field(default_factory=list)
    ignored: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    pages: list[int] = field(default_factory=list)


@dataclass(slots=True)
class ParseOutcome:
    adapter: str
    status: str = "ok"  # ok | unsupported | failed
    orders: list[ParsedOrder] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)


@dataclass(slots=True)
class PdfText:
    """Pages extracted once per file and shared by every adapter that inspects it."""

    pages: list[Any]  # imports.PdfPageData

    @property
    def first_page_text(self) -> str:
        return self.pages[0].text if self.pages else ""

    @property
    def has_text(self) -> bool:
        return any((page.text or "").strip() or page.tables for page in self.pages)


@dataclass(slots=True)
class ImportContext:
    filename: str
    content: bytes
    config: Any  # CompanyConfig
    _pdf: PdfText | None = None
    _pdf_error: str | None = None

    @property
    def suffix(self) -> str:
        return Path(self.filename).suffix.lower()

    @property
    def is_pdf(self) -> bool:
        return self.suffix == ".pdf" or self.content[:5] == b"%PDF-"

    def pdf(self) -> PdfText | None:
        """Extract text/tables once; return None for unreadable PDFs (error kept in pdf_error)."""

        if not self.is_pdf:
            return None
        if self._pdf is None and self._pdf_error is None:
            from ..imports import _pdf_pages

            try:
                self._pdf = PdfText(_pdf_pages(self.content))
            except Exception as exc:  # pdfminer raises many types for damaged files
                self._pdf_error = (
                    f"The PDF could not be read ({exc.__class__.__name__}). It may be damaged or encrypted."
                )
        return self._pdf

    @property
    def pdf_error(self) -> str | None:
        self.pdf()
        return self._pdf_error


class DocumentAdapter(Protocol):
    name: str
    label: str
    description: str

    def score(self, context: ImportContext) -> float:
        """0 = cannot read this file, 1 = certain. Must be cheap and side-effect free."""

    def parse(self, context: ImportContext) -> ParseOutcome: ...


def parse_address(text: str) -> dict[str, str]:
    """Best-effort US address from a free-text block. Missing parts stay empty."""

    lines = [re.sub(r"\s+", " ", line).strip() for line in str(text or "").splitlines() if line.strip()]
    address: dict[str, str] = {}
    for index, line in enumerate(lines):
        match = CITY_STATE_ZIP.search(line)
        if not match:
            continue
        address["city"] = match.group("city").strip(" ,")
        state = match.group("state").upper().replace(".", "")
        address["state"] = STATE_NAMES.get(state, state[:2])
        address["postal_code"] = match.group("zip")
        before = [value for value in lines[:index] if value]
        street_candidates = [value for value in before if re.match(r"^\d", value)]
        if street_candidates:
            address["street"] = street_candidates[-1]
        names = [value for value in before if value not in street_candidates]
        if names:
            address["name"] = names[0]
        break
    return address


def xlsx_bytes_to_grids(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def canonical_reference(value: str, bare_number_prefix: str = "") -> str | None:
    """Normalise an order reference: 'so 77767' -> 'SO77767', '68606' -> '<prefix>68606'."""

    text = re.sub(r"[^A-Z0-9]", "", str(value or "").upper())
    if not text:
        return None
    if text.isdigit():
        return f"{bare_number_prefix.upper()}{text}" if len(text) >= 3 else None
    return text


REFERENCE_TOKEN = re.compile(
    r"(?P<prefix>\b(?:SO|INV|PO))?[\s#-]*(?<![\d.-])(?P<digits>\d{4,7})(?![\d.-])(?P<unit>\s*(?:lbs?|in|ft|\"|pallets?))?",
    re.IGNORECASE,
)
_STATE_WORDS = "|".join(sorted([re.escape(name) for name in STATE_NAMES] + list(STATE_NAMES.values()) + ["DC"],
                                key=len, reverse=True))
_STATE_BEFORE = re.compile(rf"(?:\b(?:{_STATE_WORDS})\.?,?\s*)$", re.IGNORECASE)


def references_in_text(text: str, bare_number_prefix: str = "") -> list[str]:
    """Order references typed into a quote or shipment name.

    'SO77767 SO77769 ASMBLD' -> SO77767, SO77769; '68323 BO 66878' -> INV68323, INV66878 (with the
    prefix INV). Bare numbers count only when they have 5–6 digits and are not a ZIP code after a
    state, a quantity with a unit, or part of a date, so 'Jones Santa Fe NM 87501' and
    'Smith 2024-10-05' produce nothing.
    """

    source = str(text or "")
    found: list[str] = []
    for match in REFERENCE_TOKEN.finditer(source):
        prefix, digits = (match.group("prefix") or "").upper(), match.group("digits")
        if prefix == "PO" or match.group("unit"):
            continue
        if not prefix:
            if not 5 <= len(digits) <= 6:
                continue
            if _STATE_BEFORE.search(source[:match.start("digits")]):
                continue
        reference = canonical_reference(f"{prefix}{digits}", bare_number_prefix)
        if reference and reference not in found:
            found.append(reference)
    return found
