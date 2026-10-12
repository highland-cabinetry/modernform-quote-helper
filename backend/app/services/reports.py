from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..models import ItemMaster, PackingResult

NAVY = "17324D"
TEAL = "2F6F73"
AMBER = "D09B3C"
PALE = "EAF0F2"


def _autosize(worksheet, maximum: int = 42) -> None:
    for column_cells in worksheet.columns:
        width = min(maximum, max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: Iterable[Iterable]) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(list(row))
    _autosize(sheet)


def packing_result_xlsx(result: PackingResult) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(
        workbook,
        "Broker measurements",
        ["Pallet", "Length (in)", "Width (in)", "Height (in)", "Gross weight (lb)", "Density (pcf)", "Density class"],
        [
            [row[k] for k in ("pallet", "length", "width", "height", "weight", "density_pcf", "freight_class")]
            for row in result.quote_rows
        ],
    )
    _write_sheet(
        workbook, "Contained items",
        ["Unit", "SKU", "Host unit", "Placement", "Weight included in host (lb)", "Source"],
        [(i.unit_id, i.sku, i.host_unit_id, "inside B18", i.weight, "; ".join(i.source_refs))
         for i in result.contained_items],
    )
    summary = result.summary
    _write_sheet(
        workbook,
        "Shipment Summary",
        ["Metric", "Value"],
        [
            ("Shipment", result.shipment_name),
            ("Created", result.created_at.isoformat()),
            ("Engine", f"{result.engine} {result.engine_version}"),
            ("Optimizer", "Dedicated truck" if result.optimizer == "truckload" else "Standard LTL pallets"),
            *((("Truck floor", f"{result.truck_load['linear_feet']} ft · {result.truck_load['fits']}"),)
              if result.truck_load else ()),
            ("Total item units", summary.total_items),
            ("Packed units", summary.packed_items),
            ("Unpacked units", summary.unpacked_items),
            ("Known item weight (lb)", summary.total_item_weight),
            ("Total freight weight (lb)", summary.total_freight_weight),
            ("Pallet count", summary.total_pallets),
            ("Human review required", "YES" if result.human_review_required else "No"),
        ],
    )
    _write_sheet(
        workbook,
        "Pallets",
        [
            "Pallet",
            "Type",
            "Length",
            "Width",
            "Freight length",
            "Freight width",
            "Freight height",
            "Stack height",
            "Item weight",
            "Tare",
            "Total weight",
            "Max weight",
            "Used volume %",
            "Layer coverage %",
            "Largest cabinet-layer void %",
            "Top support coverage %",
            "Density (lb/ft³)",
            "Estimated density class",
            "Overhang by side",
            "Required actions",
            "Item quantities",
            "Warnings",
        ],
        [
            (
                pallet.pallet_number,
                pallet.pallet_type,
                pallet.length,
                pallet.width,
                pallet.load_length or pallet.length,
                pallet.load_width or pallet.width,
                pallet.freight_height,
                pallet.stack_height,
                pallet.item_weight,
                pallet.tare_weight,
                pallet.total_weight,
                pallet.max_total_weight,
                pallet.used_volume_percent,
                pallet.layer_coverage_percent,
                pallet.largest_void_percent,
                pallet.top_support_coverage_percent,
                pallet.density_pcf,
                pallet.estimated_freight_class,
                json.dumps(pallet.overhang),
                "; ".join(pallet.required_actions),
                "; ".join(f"{sku} × {quantity}" for sku, quantity in pallet.item_quantities.items()),
                "; ".join(warning.message for warning in pallet.warnings),
            )
            for pallet in result.pallets
        ],
    )
    _write_sheet(
        workbook,
        "Layer Metrics",
        [
            "Pallet",
            "Layer Z",
            "Items",
            "Cabinets",
            "Target cabinets",
            "Layer close reason",
            "Coverage %",
            "Largest void %",
            "Fragmentation %",
        ],
        [
            (
                pallet.pallet_number,
                metric.z,
                metric.item_count,
                metric.cabinet_count,
                metric.target_cabinet_count,
                metric.close_reason,
                metric.coverage_percent,
                metric.largest_void_percent,
                metric.fragmentation_percent,
            )
            for pallet in result.pallets
            for metric in pallet.layer_metrics
        ],
    )
    _write_sheet(
        workbook,
        "Override Outcomes",
        ["Unit ID", "Kind", "Status", "Requested rotation", "Applied rotation", "Message"],
        [
            (
                outcome.unit_id,
                outcome.kind,
                outcome.status,
                outcome.requested_rotation,
                outcome.applied_rotation,
                outcome.message,
            )
            for outcome in result.override_outcomes
        ],
    )
    _write_sheet(
        workbook,
        "Consolidation",
        ["Source pallets", "Target type", "Status", "Blocker", "Message"],
        [
            (
                ", ".join(str(number) for number in attempt.source_pallets),
                attempt.target_pallet_type,
                attempt.status,
                attempt.blocker_code,
                attempt.message,
            )
            for attempt in result.consolidation_attempts
        ],
    )
    _write_sheet(
        workbook,
        "3D Placements",
        [
            "Pallet",
            "Unit ID",
            "Original SKU",
            "Master SKU",
            "Packing state",
            "Description",
            "Category",
            "Item kind",
            "Source",
            "X",
            "Y",
            "Z",
            "Length",
            "Width",
            "Height",
            "Flatpack master dimensions",
            "Assembled master dimensions",
            "Rotation",
            "Weight",
            "Support %",
            "Packing role",
            "Bundle ID",
            "Bundle position",
            "Bundle size",
            "Bundle approved",
            "Pair group ID",
            "Paired family",
            "Paired unit IDs",
            "Dunnage required",
            "Support contacts",
            "Locked",
            "Overhang approved",
        ],
        [
            (
                pallet.pallet_number,
                item.unit_id,
                item.original_sku,
                item.sku,
                item.packing_state.value,
                item.description,
                item.category,
                item.item_kind.value,
                "; ".join(item.source_refs),
                item.x,
                item.y,
                item.z,
                item.length,
                item.width,
                item.height,
                " × ".join(str(value) for value in item.flatpack_dimensions or ()),
                " × ".join(str(value) for value in item.assembled_dimensions or ()),
                item.rotation,
                item.weight,
                round(item.support_ratio * 100, 1),
                item.packing_role.value,
                item.bundle_id,
                item.bundle_position,
                item.bundle_size,
                "YES" if item.bundle_approved else "No",
                item.pair_group_id,
                item.paired_family,
                "; ".join(item.paired_unit_ids),
                "YES" if item.requires_dunnage else "No",
                json.dumps([contact.model_dump(mode="json") for contact in item.support_contacts]),
                "YES" if item.locked else "No",
                "YES" if item.overhang_approved else "No",
            )
            for pallet in result.pallets
            for item in pallet.items
        ],
    )
    _write_sheet(
        workbook,
        "Unpacked Review",
        [
            "Unit ID",
            "Original SKU",
            "Master SKU",
            "Packing state",
            "Description",
            "Source",
            "Weight",
            "Dimensions",
            "Code",
            "Reason",
        ],
        [
            (
                item.unit_id,
                item.original_sku,
                item.sku,
                item.packing_state.value,
                item.description,
                "; ".join(item.source_refs),
                item.weight,
                " × ".join(str(value) for value in item.dimensions),
                item.reason_code,
                item.reason,
            )
            for item in result.unpacked
        ],
    )
    _write_sheet(
        workbook,
        "Warnings",
        ["Severity", "Code", "Pallet", "Message", "Unit IDs"],
        [
            (issue.severity.value, issue.code, issue.pallet_number, issue.message, "; ".join(issue.unit_ids))
            for issue in result.warnings
        ],
    )
    if result.shipping_info:
        info = result.shipping_info
        _write_sheet(
            workbook,
            "Quote Information",
            ["Field", "Value"],
            [
                ("Confirmed", info.confirmed),
                ("References", ", ".join(info.customer_references)),
                ("Origin", json.dumps(info.origin.model_dump(mode="json"))),
                ("Destination", json.dumps(info.destination.model_dump(mode="json"))),
                ("Commodity", info.commodity_description),
                ("NMFC", info.nmfc),
                ("Accessorials", json.dumps(info.accessorials.model_dump(mode="json"))),
            ],
        )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def packing_result_csv(result: PackingResult) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(
        [
            "status",
            "pallet",
            "pallet_type",
            "unit_id",
            "original_sku",
            "master_sku",
            "packing_state",
            "description",
            "item_kind",
            "source",
            "x",
            "y",
            "z",
            "length",
            "width",
            "height",
            "rotation",
            "weight",
            "packing_role",
            "bundle_id",
            "bundle_position",
            "bundle_size",
            "bundle_approved",
            "pair_group_id",
            "paired_family",
            "paired_unit_ids",
            "requires_dunnage",
            "reason_code",
            "reason",
        ]
    )
    for pallet in result.pallets:
        for item in pallet.items:
            writer.writerow(
                [
                    "packed",
                    pallet.pallet_number,
                    pallet.pallet_type,
                    item.unit_id,
                    item.original_sku,
                    item.sku,
                    item.packing_state.value,
                    item.description,
                    item.item_kind.value,
                    "; ".join(item.source_refs),
                    item.x,
                    item.y,
                    item.z,
                    item.length,
                    item.width,
                    item.height,
                    item.rotation,
                    item.weight,
                    item.packing_role.value,
                    item.bundle_id or "",
                    item.bundle_position or "",
                    item.bundle_size or "",
                    item.bundle_approved,
                    item.pair_group_id or "",
                    item.paired_family or "",
                    "; ".join(item.paired_unit_ids),
                    item.requires_dunnage,
                    "",
                    "",
                ]
            )
    for item in result.contained_items:
        pallet = next(p for p in result.pallets if any(host.unit_id == item.host_unit_id for host in p.items))
        row = [""] * 29
        row[:10] = ["contained", pallet.pallet_number, pallet.pallet_type, item.unit_id, item.original_sku,
            item.sku, item.packing_state.value, "inside B18", "accessory", "; ".join(item.source_refs)]
        row[17] = 0  # host row includes the contained weight; do not double-sum it
        row[28] = f"inside B18 {item.host_unit_id}; {item.weight:g} lb already included in host weight"
        writer.writerow(row)
    for item in result.unpacked:
        writer.writerow(
            [
                "unpacked",
                "",
                "",
                item.unit_id,
                item.original_sku,
                item.sku,
                item.packing_state.value,
                item.description,
                "",
                "; ".join(item.source_refs),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                *item.dimensions,
                "",
                item.weight,
                "",
                "",
                "",
                "",
                "",
                item.reason_code,
                item.reason,
            ]
        )
    return stream.getvalue().encode("utf-8-sig")


def packing_result_pdf(result: PackingResult) -> bytes:
    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=letter,
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
        title=f"Pallet report - {result.shipment_name}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=7.5, leading=9, alignment=TA_LEFT))
    styles.add(
        ParagraphStyle(
            name="Eyebrow",
            parent=styles["BodyText"],
            fontSize=7,
            leading=8,
            textColor=colors.HexColor(f"#{TEAL}"),
            spaceAfter=3,
        )
    )
    story = [
        Paragraph("FREIGHT QUOTE HELPER · PALLET REPORT", styles["Eyebrow"]),
        Paragraph(result.shipment_name, styles["Title"]),
        Paragraph(result.created_at.strftime("Generated %Y-%m-%d %H:%M UTC"), styles["Small"]),
        Spacer(1, 0.12 * inch),
    ]
    summary_data = [
        ["Items", "Known item wt.", "Pallets", "Freight wt.", "Review"],
        [
            str(result.summary.total_items),
            f"{result.summary.total_item_weight:,.1f} lb",
            str(result.summary.total_pallets),
            f"{result.summary.total_freight_weight:,.1f} lb",
            "REQUIRED" if result.human_review_required else "Clear",
        ],
    ]
    summary_table = Table(summary_data, colWidths=[1.15 * inch, 1.45 * inch, 1.05 * inch, 1.45 * inch, 1.25 * inch])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor(f"#{PALE}")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AAB7BE")),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([summary_table, Spacer(1, 0.18 * inch)])

    if result.consolidation_attempts:
        story.append(Paragraph("Pallet consolidation", styles["Heading2"]))
        for attempt in result.consolidation_attempts:
            source = ", ".join(str(number) for number in attempt.source_pallets)
            blocker = f" · {attempt.blocker_code}" if attempt.blocker_code else ""
            story.append(
                Paragraph(
                    f"<b>{attempt.status.upper()}{blocker}</b> — pallets {source}: {attempt.message}",
                    styles["Small"],
                )
            )
        story.append(Spacer(1, 0.12 * inch))

    if result.warnings:
        story.append(Paragraph("Warnings & required review", styles["Heading2"]))
        for issue in result.warnings:
            location = f" · pallet {issue.pallet_number}" if issue.pallet_number else ""
            story.append(
                Paragraph(
                    f"<b>{issue.severity.value.upper()} · {issue.code}{location}</b> — {issue.message}", styles["Small"]
                )
            )
        story.append(Spacer(1, 0.12 * inch))

    if result.contained_items:
        from xml.sax.saxutils import escape
        story.append(Paragraph("Contained items (weight included in host)", styles["Heading2"]))
        for item in result.contained_items:
            text = f"{item.sku} ({item.unit_id}) inside B18 {item.host_unit_id}: {item.weight:g} lb"
            story.append(Paragraph(escape(text), styles["Small"]))
        story.append(Spacer(1, 0.12 * inch))

    story.append(Paragraph("Measurements sent to brokers (rounded inches / pounds)", styles["Heading2"]))
    for row in result.quote_rows:
        story.append(
            Paragraph(
                f"P{row['pallet']}: {row['dimensions']} · {row['weight']} lb · "
                f"{row['density_pcf']} pcf · class {row['freight_class']}",
                styles["Small"],
            )
        )

    for pallet in result.pallets:
        story.append(Paragraph(f"Pallet {pallet.pallet_number} · {pallet.display_name}", styles["Heading2"]))
        story.append(
            Paragraph(
                f"Freight dimensions: {pallet.load_length or pallet.length:.1f} × "
                f"{pallet.load_width or pallet.width:.1f} × "
                f"{pallet.freight_height:.1f} in &nbsp; | &nbsp; "
                f"Weight: {pallet.total_weight:,.1f} lb &nbsp; | &nbsp; "
                f"Density class: {pallet.estimated_freight_class} ({pallet.density_pcf:.2f} lb/ft³) &nbsp; | &nbsp; "
                f"Layer coverage: {pallet.layer_coverage_percent:.1f}% &nbsp; | &nbsp; "
                f"Largest void: {pallet.largest_void_percent:.1f}% &nbsp; | &nbsp; "
                f"Top support: {pallet.top_support_coverage_percent:.1f}% &nbsp; | &nbsp; "
                f"Used volume: {pallet.used_volume_percent:.1f}%",
                styles["Small"],
            )
        )
        for action in pallet.required_actions:
            story.append(Paragraph(f"<b>Warehouse action:</b> {action}", styles["Small"]))
        rows = [["SKU", "State", "Kind / relationship", "Description", "Qty", "Weight", "Sources"]]
        quantity_keys = sorted({(item.original_sku, item.packing_state.value) for item in pallet.items})
        for sku, state in quantity_keys:
            matching = [item for item in pallet.items if item.original_sku == sku and item.packing_state.value == state]
            rows.append(
                [
                    sku,
                    state,
                    Paragraph(
                        " · ".join(
                            filter(
                                None,
                                [
                                    matching[0].item_kind.value.replace("_", " ") if matching else "",
                                    f"bundle {matching[0].bundle_size}" if matching and matching[0].bundle_id else "",
                                    f"paired {matching[0].paired_family}"
                                    if matching and matching[0].pair_group_id
                                    else "",
                                ],
                            )
                        ),
                        styles["Small"],
                    ),
                    Paragraph(matching[0].description if matching else "", styles["Small"]),
                    len(matching),
                    f"{sum(item.weight for item in matching):,.1f} lb",
                    Paragraph(
                        "; ".join(sorted({source for item in matching for source in item.source_refs})), styles["Small"]
                    ),
                ]
            )
        table = Table(
            rows,
            repeatRows=1,
            colWidths=[0.72 * inch, 0.52 * inch, 0.95 * inch, 1.45 * inch, 0.35 * inch, 0.65 * inch, 1.55 * inch],
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TEAL}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B7C3C7")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7F8")]),
                ]
            )
        )
        story.extend([table, Spacer(1, 0.16 * inch)])

    if result.unpacked:
        story.extend([PageBreak(), Paragraph("Unpacked / manual review", styles["Heading1"])])
        rows = [["SKU", "State", "Source", "Dimensions", "Weight", "Reason"]]
        for item in result.unpacked:
            rows.append(
                [
                    item.original_sku,
                    item.packing_state.value,
                    Paragraph("; ".join(item.source_refs), styles["Small"]),
                    " × ".join(f"{value:g}" for value in item.dimensions),
                    f"{item.weight:g} lb",
                    Paragraph(item.reason, styles["Small"]),
                ]
            )
        table = Table(
            rows,
            repeatRows=1,
            colWidths=[0.8 * inch, 0.65 * inch, 1.05 * inch, 0.75 * inch, 0.6 * inch, 2.55 * inch],
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#A94F45")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)

    story.extend([Spacer(1, 0.2 * inch), Paragraph(result.disclaimer, styles["Small"])])
    document.build(story)
    return stream.getvalue()


def item_master_template_xlsx(items: list[ItemMaster] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Item Master"
    headers = list(ItemMaster.model_fields.keys())
    headers.remove("id")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    for item in items or []:
        payload = item.model_dump(mode="json")
        sheet.append([",".join(payload[name]) if name == "allowed_rotations" else payload[name] for name in headers])
    _autosize(sheet)
    prefix_sheet = workbook.create_sheet("SKU Prefixes")
    prefix_sheet.append(["Cabinet Color", "Item Prefix"])
    prefixes = [
        ("Espresso Shaker", "ESS-"),
        ("Gray Shaker", "GRS-"),
        ("Modern Slate Shaker", "MSL-"),
        ("Mountain Timber Shaker", "MTS-"),
        ("Onyx Black Shaker", "ONB-"),
        ("Signature Blue Shaker", "SBS-"),
        ("Slim Dove White Shaker", "SDW-"),
        ("Slim Saige Green Shaker", "SSG-"),
        ("White Shaker", "WHS-"),
    ]
    for color, prefix in prefixes:
        prefix_sheet.append([color, prefix])
    for cell in prefix_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    _autosize(prefix_sheet)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


# The simple catalog workbook (same layout as the original Quote Helper "Shaker item list"):
# a "Cabinet Colors" sheet (color name → item prefix) and an "Item References" sheet with one row per
# product. Handling rules (stacking, fragile, top-load ...) are kept in the app and are not part of it;
# importing this workbook changes only the columns it contains.
SIMPLE_CATALOG_COLUMNS: tuple[tuple[str, str], ...] = (
    ("In House Item SKU", "sku"),
    ("Manufacturer Item SKU", "manufacturer_sku"),
    ("Manufacturer Dimensions", "manufacturer_dimensions"),
    ("Unit Weight (lb)", "weight"),
    ("Flatpack L", "flatpack_length"),
    ("Flatpack W", "flatpack_width"),
    ("Flatpack H", "flatpack_height"),
    ("Assembled L", "assembled_length"),
    ("Assembled W", "assembled_width"),
    ("Assembled H", "assembled_height"),
    ("Category", "category"),
    ("Description", "description"),
)
COLOR_SHEET = "Cabinet Colors"
ITEM_SHEET = "Item References"


def _simple_value(item: ItemMaster, field: str):
    # Full precision (the sheet displays two decimals), so exporting and importing changes nothing.
    return getattr(item, field)


def simple_catalog_xlsx(items: list[ItemMaster], colors: Iterable = ()) -> bytes:
    workbook = Workbook()
    color_sheet = workbook.active
    color_sheet.title = COLOR_SHEET
    color_sheet.append(["Cabinet Color", "Item Prefix"])
    for color in colors:
        color_sheet.append([color.name, color.prefix])
    sheet = workbook.create_sheet(ITEM_SHEET)
    sheet.append([label for label, _ in SIMPLE_CATALOG_COLUMNS])
    for item in sorted(items, key=lambda value: (value.category or "", value.sku.upper())):
        sheet.append([_simple_value(item, field) for _, field in SIMPLE_CATALOG_COLUMNS])
    for worksheet in (color_sheet, sheet):
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color="FFFFFF", bold=True)
        worksheet.freeze_panes = "A2"
        _autosize(worksheet)
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=10):
        for cell in row:
            cell.number_format = "0.00"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def simple_catalog_csv(items: list[ItemMaster]) -> bytes:
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow([label for label, _ in SIMPLE_CATALOG_COLUMNS])
    for item in sorted(items, key=lambda value: (value.category or "", value.sku.upper())):
        writer.writerow(["" if (value := _simple_value(item, field)) is None else value
                         for _, field in SIMPLE_CATALOG_COLUMNS])
    return stream.getvalue().encode("utf-8-sig")

def packing_result_json(result: PackingResult) -> bytes:
    return json.dumps(result.model_dump(mode="json", exclude={"broker_rates"}), indent=2).encode("utf-8")
