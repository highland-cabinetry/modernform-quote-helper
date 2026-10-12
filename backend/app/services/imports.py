from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..models import ItemMaster, PackingSlipLine, PackingState

COLUMN_ALIASES = {
    "sku": {
        "sku", "item", "itemnumber", "itemno", "partnumber", "partno", "productcode", "code",
        "inhouseitemsku", "itemsku", "productid", "productidentifier",
    },
    "manufacturer_sku": {"manufacturersku", "manufactureritemsku", "mfrsku"},
    "description": {"description", "desc", "itemdescription", "product"},
    "manufacturer_dimensions": {"manufacturerdimensions", "mfrdimensions"},
    "quantity_ordered": {"quantity", "qty", "ordered", "orderqty", "qtyordered", "orderedqty"},
    "quantity_shipped": {"shipped", "shipqty", "qtyshipped", "shippedqty", "quantityshipped"},
    "length": {"length", "len", "l", "depth"},
    "width": {"width", "wid", "w"},
    "height": {"height", "ht", "h"},
    "flatpack_length": {"flatpacklength", "flatpackl", "universall"},
    "flatpack_width": {"flatpackwidth", "flatpackw", "universalw"},
    "flatpack_height": {"flatpackheight", "flatpackh"},
    "assembled_length": {"assembledlength", "assembledl"},
    "assembled_width": {"assembledwidth", "assembledw"},
    "assembled_height": {"assembledheight", "assembledh"},
    "weight": {"weight", "wt", "lbs", "unitweight", "unitweightlb"},
    "packing_state": {"packingstate", "assemblystate", "assembly", "assemble", "assembled", "packstate"},
    "category": {"category", "type", "productcategory"},
    "item_kind": {"itemkind", "physicalitemkind", "packingitemkind"},
    "stackable": {"stackable", "canstack"},
    "fragile": {"fragile"},
    "must_stay_upright": {"muststayupright", "upright", "keepupright"},
    "no_upside_down": {"noupsidedown", "thissideup"},
    "allowed_rotations": {"allowedrotations", "rotations"},
    "max_stack_layers": {"maxstacklayers", "stacklayers", "maxlayers"},
    "max_weight_on_top": {"maxweightontop", "topweight"},
    "load_bearing_strength": {"loadbearingstrength", "loadbearing", "bearing"},
    "preferred_pallet_type": {"preferredpallettype", "pallettype", "preferredpallet"},
    "freight_class": {"freightclass", "class"},
    "nmfc": {"nmfc", "nmfcnumber"},
    "color": {"color", "hexcolor"},
    "rules_reviewed": {"rulesreviewed", "handlingrulesreviewed"},
    "packing_role": {"packingrole", "warehouserole", "palletizingrole"},
    "bundle_max_units": {"bundlemaxunits", "bundlequantity", "maxbundleunits"},
    "bundle_approved": {"bundleapproved", "approvedforbundling", "canbundle"},
    "allow_edge_standing": {"allowedgestanding", "canstandonedge", "edgestanding"},
    "overhang_approved": {"overhangapproved", "approvedforoverhang", "canoverhang"},
    "max_overhang_per_side": {"maxoverhangperside", "accessoryoverhang", "maxoverhang"},
    "active": {"active", "enabled"},
}


@dataclass(slots=True)
class ItemMasterImportResult:
    items: list[ItemMaster]
    skipped: list[dict[str, object]]
    # For each item, the ItemMaster fields its file actually has a column for. A simple catalog
    # workbook (SKU, sizes, weight, category, description) then updates only those fields and keeps
    # an existing product's handling rules.
    provided: list[set[str]] | None = None


RECORD_FIELDS = {"length": "flatpack_length", "width": "flatpack_width", "height": "flatpack_height"}
DIMENSION_FIELDS = ("flatpack_length", "flatpack_width", "flatpack_height",
                    "assembled_length", "assembled_width", "assembled_height")


@dataclass(slots=True)
class PackingSlipImportResult:
    lines: list[PackingSlipLine]
    ignored: list[dict[str, object]]
    shipping_suggestions: list[dict[str, object]]
    documents: list[dict[str, object]]
    suggested_shipment_name: str | None = None


@dataclass(slots=True)
class PdfPageData:
    page_number: int
    text: str
    tables: list[list[list[Any]]]
    document_index: int = 0
    packing_state: PackingState = PackingState.FLATPACK


NON_PHYSICAL_ITEM_PREFIXES = (
    "assembly",
    "discount",
    "freight",
    "handling",
    "labor",
    "othercharge",
    "shipping",
    "subtotal",
    "tax",
    "total",
    "salestax",
    "surcharge",
    "fuelsurcharge",
    "deposit",
    "payment",
    "restocking",
)

COLOR_SECTION_SKUS = {"ESS", "GRS", "MSL", "MTS", "ONB", "SBS", "SDW", "SSG", "WHS"}


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _canonical_headers(headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, raw in enumerate(headers):
        normalized = _header_key(raw)
        for canonical, aliases in COLUMN_ALIASES.items():
            if normalized in aliases:
                mapping[index] = canonical
                break
    return mapping


def _records_from_grid(grid: list[list[Any]]) -> list[dict[str, Any]]:
    return [record for _, record in _records_with_rows(grid)]


def _records_with_rows(grid: list[list[Any]]) -> list[tuple[int, dict[str, Any]]]:
    if not grid:
        return []
    header_index = 0
    header_map: dict[int, str] = {}
    for candidate_index, row in enumerate(grid[:20]):
        candidate = _canonical_headers(row)
        quantity_fields = {"quantity_ordered", "quantity_shipped"}
        if "sku" in candidate.values() and (quantity_fields.intersection(candidate.values()) or len(candidate) >= 2):
            header_index = candidate_index
            header_map = candidate
            break
    if not header_map:
        return []
    records: list[tuple[int, dict[str, Any]]] = []
    for row_number, row in enumerate(grid[header_index + 1 :], start=header_index + 2):
        record = {canonical: row[index] if index < len(row) else None for index, canonical in header_map.items()}
        if str(record.get("sku") or "").strip():
            records.append((row_number, record))
    return records


def _parse_number(value: Any, default: float | None = None) -> float | None:
    if value is None or value == "":
        return default
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    try:
        return float(cleaned)
    except (TypeError, ValueError):
        return default


def _parse_bool(value: Any, default: bool) -> bool:
    if value is None or value == "":
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y", "x"}


def _parse_rotations(value: Any, *, upright: bool) -> list[str]:
    if value:
        parsed = [token.strip().upper() for token in re.split(r"[,;|\s]+", str(value)) if token.strip()]
        if parsed:
            return parsed
    return ["LWH", "WLH"] if upright else ["LWH", "WLH", "LHW", "HLW", "WHL", "HWL"]


def _parse_packing_state(value: Any) -> PackingState:
    normalized = _header_key(value)
    if normalized in {"assembled", "assembly", "assemble", "yes", "y", "true", "1", "x"} or \
            str(value).strip() in {"✓", "✔"}:
        return PackingState.ASSEMBLED
    return PackingState.FLATPACK


def _selected_quantity(record: dict[str, Any]) -> Any:
    """Use shipped quantity when authored; otherwise use the ordered quantity."""

    shipped = record.get("quantity_shipped")
    if shipped is not None and str(shipped).strip() != "":
        return shipped
    return record.get("quantity_ordered")


def _quantity_error(value: Any) -> tuple[int | None, tuple[str, str] | None]:
    if value is None or str(value).strip() == "":
        return None, ("MISSING_QUANTITY", "No quantity was supplied; headings and non-item rows are not imported.")
    text = str(value).strip().replace(",", "")
    try:
        parsed = float(text)
    except (TypeError, ValueError):
        return None, ("INVALID_QUANTITY", "Quantity must be a positive whole number.")
    if not parsed.is_integer() or parsed <= 0 or parsed > 10_000:
        return None, ("INVALID_QUANTITY", "Quantity must be a positive whole number between 1 and 10,000.")
    return int(parsed), None


def _non_physical_reason(sku: str, description: str = "") -> tuple[str, str] | None:
    normalized = _header_key(sku)
    if normalized.startswith(NON_PHYSICAL_ITEM_PREFIXES):
        return "NON_PHYSICAL_ROW", "Service, fee, discount, tax, or summary row; no physical item was imported."
    if sku.strip().upper() in COLOR_SECTION_SKUS and re.search(r"\*{3,}", description):
        return "SECTION_HEADING", "Cabinet-color section label; no physical item was imported."
    return None


def _csv_grid(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    # Only the delimiter is sniffed. The Sniffer's quote/doublequote guesses depend on which rows
    # land in the sample and can misread Excel's "" escaping (42"" -> 42""), so quoting stays standard.
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        delimiter = ","
    return [list(row) for row in csv.reader(io.StringIO(text), csv.excel, delimiter=delimiter)]


def _xlsx_grids(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]


def _pdf_pages(content: bytes) -> list[PdfPageData]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - installation-dependent guard
        raise ValueError("PDF support requires the optional pdfplumber dependency") from exc

    pages: list[PdfPageData] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            pages.append(
                PdfPageData(
                    page_number=page_number,
                    text=page.extract_text() or "",
                    tables=[table for table in (page.extract_tables() or []) if table],
                )
            )
    return pages


def _table_value(tables: list[list[list[Any]]], header: str) -> str:
    target = _header_key(header)
    for table in tables:
        for row_index, row in enumerate(table[:-1]):
            for column_index, cell in enumerate(row):
                if _header_key(cell) == target:
                    next_row = table[row_index + 1]
                    if column_index < len(next_row):
                        return str(next_row[column_index] or "").strip()
    return ""


def _document_type(text: str) -> str:
    heading = text[:800]
    if re.search(r"\bpacking\s+slip\b", heading, re.I):
        return "packing_slip"
    if re.search(r"\binvoice\b", heading, re.I):
        return "invoice"
    if re.search(r"\bsales\s+order\b", heading, re.I):
        return "sales_order"
    return "continuation"


def _page_identifiers(page: PdfPageData) -> tuple[str | None, str | None]:
    invoice_value = _table_value(page.tables, "Invoice #")
    invoice_digits = re.search(r"\d{3,}", invoice_value)
    invoice_number = invoice_digits.group(0) if invoice_digits else None
    sales_order_match = re.search(r"\bSO[\s#-]*(\d{3,})\b", page.text, re.I)
    sales_order_number = f"SO{sales_order_match.group(1)}" if sales_order_match else None
    return invoice_number, sales_order_number


def _page_packing_state(page: PdfPageData) -> PackingState:
    assembly = _table_value(page.tables, "Assembly")
    if assembly:
        return _parse_packing_state(assembly)
    if re.search(r"\*+\s*assembled\s*\*+|\bassemble(?:d)?\s+freight\b", page.text, re.I):
        return PackingState.ASSEMBLED
    return PackingState.FLATPACK


def _ship_to_address(page: PdfPageData) -> dict[str, str]:
    for table in page.tables:
        if not table or not table[0] or _header_key(table[0][0]) != "shipto":
            continue
        value = "\n".join(str(cell or "") for row in table[1:] for cell in row if cell)
        locality = re.search(r"([^\n,]+),\s*([A-Z]{2})\s+(\d{5})(?:-\d{4})?", value, re.I)
        postal = locality or re.search(r"\b(\d{5})(?:-\d{4})?\b", value)
        address: dict[str, str] = {}
        if postal:
            address["postal_code"] = postal.group(3) if locality else postal.group(1)
        if locality:
            address["city"] = locality.group(1).strip()
            address["state"] = locality.group(2).upper()
        return address
    return {}


def _page_accessorials(text: str) -> tuple[dict[str, bool], dict[str, bool]]:
    residential_negative = bool(re.search(r"\b(?:no|without|non[- ]?)\s*residential\b", text, re.I))
    liftgate_negative = bool(re.search(r"\b(?:no|without)\s+lift\s*gate\b", text, re.I))
    positive_text = re.sub(r"\b(?:no|without)\s+lift\s*gate\b", "", text, flags=re.I)
    positive_text = re.sub(r"\b(?:no|without|non[- ]?)\s*residential\b", "", positive_text, flags=re.I)
    positive = {
        "residential": bool(re.search(r"\b(?:residential|job\s*site|jobsite|home delivery)\b", positive_text, re.I)),
        "liftgate": bool(re.search(r"\blift\s*gate\b", positive_text, re.I)),
        "appointment_delivery": bool(re.search(r"\bappointment\b", positive_text, re.I)),
        "delivery_call_ahead": bool(re.search(r"\bcall\s*(?:ahead|before)\b", positive_text, re.I)),
    }
    negative = {"residential": residential_negative, "liftgate": liftgate_negative}
    return positive, negative


def _pdf_documents(
    filename: str,
    pages: list[PdfPageData],
) -> tuple[list[dict[str, object]], dict[str, object] | None, str | None]:
    documents: list[dict[str, object]] = []
    document_lookup: dict[tuple[str, str], int] = {}
    current_index = -1
    for page in pages:
        page_type = _document_type(page.text)
        invoice_number, sales_order_number = _page_identifiers(page)
        if page_type in {"invoice", "packing_slip"}:
            primary_reference = f"INV{invoice_number}" if invoice_number else sales_order_number
        elif page_type == "sales_order":
            primary_reference = sales_order_number
        else:
            primary_reference = None
        key = (page_type, primary_reference or "")
        if page_type != "continuation" and primary_reference and key in document_lookup:
            current_index = document_lookup[key]
        elif page_type != "continuation":
            current_index = len(documents)
            document_lookup[key] = current_index
            documents.append(
                {
                    "document_type": page_type,
                    "primary_reference": primary_reference,
                    "invoice_number": invoice_number,
                    "sales_order_number": sales_order_number,
                    "packing_state": _page_packing_state(page).value,
                    "pages": [],
                    "destination_zips": [],
                    "accessorials": {"residential": False, "liftgate": False},
                    "negative_accessorials": {"residential": False, "liftgate": False},
                    "source_provenance": [],
                    "warnings": [],
                }
            )
        elif current_index < 0:
            current_index = len(documents)
            documents.append(
                {
                    "document_type": "unknown",
                    "primary_reference": None,
                    "invoice_number": None,
                    "sales_order_number": None,
                    "packing_state": _page_packing_state(page).value,
                    "pages": [],
                    "destination_zips": [],
                    "accessorials": {"residential": False, "liftgate": False},
                    "negative_accessorials": {"residential": False, "liftgate": False},
                    "source_provenance": [],
                    "warnings": [],
                }
            )
        page.document_index = current_index
        document = documents[current_index]
        page.packing_state = PackingState(str(document["packing_state"]))
        document["pages"].append(page.page_number)  # type: ignore[union-attr]
        address = _ship_to_address(page)
        postal_code = address.get("postal_code")
        if postal_code and postal_code not in document["destination_zips"]:  # type: ignore[operator]
            document["destination_zips"].append(postal_code)  # type: ignore[union-attr]
        positive, negative = _page_accessorials(page.text)
        for name in ("residential", "liftgate"):
            document["accessorials"][name] = bool(document["accessorials"][name] or positive[name])  # type: ignore[index]
            document["negative_accessorials"][name] = bool(  # type: ignore[index]
                document["negative_accessorials"][name] or negative[name]  # type: ignore[index]
            )
        document["source_provenance"].append(f"{filename}: page {page.page_number}")  # type: ignore[union-attr]

    primary_references = [str(doc["primary_reference"]) for doc in documents if doc.get("primary_reference")]
    primary_references = list(dict.fromkeys(primary_references))
    suggested_name = " + ".join(primary_references) or None
    all_references: list[str] = []
    destination_zips: list[str] = []
    warnings: list[str] = []
    for document in documents:
        for reference in (document.get("primary_reference"), document.get("sales_order_number")):
            if reference and reference not in all_references:
                all_references.append(str(reference))
        for postal_code in document["destination_zips"]:  # type: ignore[union-attr]
            if postal_code not in destination_zips:
                destination_zips.append(str(postal_code))
        for name in ("residential", "liftgate"):
            if document["accessorials"][name] and document["negative_accessorials"][name]:  # type: ignore[index]
                message = f"Conflicting positive and negative {name} indications; positive service was selected."
                document["warnings"].append(message)  # type: ignore[union-attr]
                warnings.append(message)
    if len(destination_zips) > 1:
        warnings.append(
            "Combined documents contain different delivery ZIP codes; select the intended ZIP before broker handoff."
        )
    suggestion = None
    if all_references or destination_zips:
        suggestion = {
            "file": filename,
            "customer_references": all_references,
            "destination": {"postal_code": destination_zips[0]} if len(destination_zips) == 1 else {},
            "destination_zip_candidates": destination_zips,
            "accessorials": {
                "residential": any(bool(doc["accessorials"]["residential"]) for doc in documents),  # type: ignore[index]
                "liftgate": any(bool(doc["accessorials"]["liftgate"]) for doc in documents),  # type: ignore[index]
                "appointment_delivery": False,
                "delivery_call_ahead": False,
            },
            "confirmed": False,
            "warnings": list(dict.fromkeys(warnings)),
            "source_provenance": [
                provenance
                for document in documents
                for provenance in document["source_provenance"]  # type: ignore[union-attr]
            ],
        }
    return documents, suggestion, suggested_name


def _pdf_sources(pages: list[PdfPageData]) -> list[tuple[str, list[list[Any]], PackingState]]:
    sources: list[tuple[str, list[list[Any]], PackingState]] = []
    for page in pages:
        for table_number, table in enumerate(page.tables, start=1):
            sources.append((f"page {page.page_number}, table {table_number}", table, page.packing_state))
        if page.tables:
            continue
        rows: list[list[Any]] = [["Quantity", "SKU", "Description"]]
        pattern = re.compile(r"^\s*(\d+)\s+([A-Z0-9][A-Z0-9._/\-]{2,})\s+(.+)$", re.I)
        for line in page.text.splitlines():
            if match := pattern.match(line):
                rows.append([match.group(1), match.group(2), match.group(3)])
        if len(rows) > 1:
            sources.append((f"page {page.page_number}, text", rows, page.packing_state))
    return sources


def parse_packing_slip_with_report(filename: str, content: bytes) -> PackingSlipImportResult:
    suffix = Path(filename).suffix.lower()
    documents: list[dict[str, object]] = []
    suggested_shipment_name: str | None = None
    shipping_suggestions: list[dict[str, object]] = []
    if suffix in {".csv", ".tsv", ".txt"}:
        sources = [(filename, _csv_grid(content), PackingState.FLATPACK)]
    elif suffix in {".xlsx", ".xlsm"}:
        sources = [
            (f"{filename} · {sheet}", grid, PackingState.FLATPACK) for sheet, grid in _xlsx_grids(content)
        ]
    elif suffix == ".pdf":
        pages = _pdf_pages(content)
        documents, suggestion, suggested_shipment_name = _pdf_documents(filename, pages)
        sources = [
            (f"{filename} · {source_ref}", grid, packing_state)
            for source_ref, grid, packing_state in _pdf_sources(pages)
        ]
        if suggestion:
            shipping_suggestions.append(suggestion)
    else:
        raise ValueError("Supported packing slip types are CSV, TSV, XLSX, XLSM, and PDF")

    parsed: list[PackingSlipLine] = []
    ignored: list[dict[str, object]] = []
    for source_ref, grid, source_packing_state in sources:
        for row_number, record in _records_with_rows(grid):
            original_sku = str(record.get("sku") or "").strip()
            description = str(record.get("description") or "").strip()
            reason = _non_physical_reason(original_sku, description)
            selected_quantity = _selected_quantity(record)
            quantity, quantity_reason = _quantity_error(selected_quantity)
            reason = reason or quantity_reason
            if reason:
                ignored.append(
                    {
                        "source_ref": source_ref,
                        "row": row_number,
                        "original_sku": original_sku,
                        "quantity": selected_quantity,
                        "description": description,
                        "reason_code": reason[0],
                        "reason": reason[1],
                    }
                )
                continue
            parsed.append(
                PackingSlipLine(
                    source_ref=source_ref,
                    original_sku=original_sku,
                    quantity=quantity,
                    packing_state=(
                        _parse_packing_state(record.get("packing_state"))
                        if str(record.get("packing_state") or "").strip()
                        else source_packing_state
                    ),
                    description=description,
                    length=_parse_number(record.get("length")),
                    width=_parse_number(record.get("width")),
                    height=_parse_number(record.get("height")),
                    weight=_parse_number(record.get("weight")),
                    category=str(record.get("category") or "").strip() or None,
                    line_number=row_number,
                )
            )
    return PackingSlipImportResult(
        lines=parsed,
        ignored=ignored,
        shipping_suggestions=shipping_suggestions,
        documents=documents,
        suggested_shipment_name=suggested_shipment_name,
    )


def parse_packing_slip(filename: str, content: bytes) -> list[PackingSlipLine]:
    """Compatibility wrapper for callers that only need accepted physical lines."""

    result = parse_packing_slip_with_report(filename, content)
    if not result.lines:
        raise ValueError(
            "No physical item rows were found. Include recognizable SKU/item and positive quantity headers."
        )
    return result.lines


def parse_item_master_with_report(filename: str, content: bytes) -> ItemMasterImportResult:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        grids = [(filename, _csv_grid(content))]
    elif suffix in {".xlsx", ".xlsm"}:
        grids = _xlsx_grids(content)
    else:
        raise ValueError("Item master import supports CSV, TSV, XLSX, and XLSM")

    items: list[ItemMaster] = []
    provided: list[set[str]] = []
    skipped: list[dict[str, object]] = []
    for source_name, grid in grids:
        for row_number, record in enumerate(_records_from_grid(grid), start=2):
            try:
                for name in ("flatpack_length", "flatpack_width", "flatpack_height", "length", "width", "height",
                             "assembled_length", "assembled_width", "assembled_height"):
                    # 0 or blank means "no carton of this kind" (simple item lists use 0).
                    if (number := _parse_number(record.get(name))) is not None and number <= 0:
                        record[name] = None
                if _parse_number(record.get("weight")) is None:
                    raise ValueError("Unit Weight (lb) is missing")
                upright = _parse_bool(record.get("must_stay_upright"), True)
                category = str(record.get("category") or "Miscellaneous").strip()
                inferred_kind = (
                    "cabinet_component" if category.lower() == "face frame and door"
                    else "accessory" if category.lower() == "accessory"
                    else "cabinet" if "cabinet" in category.lower()
                    else "other"
                )
                item_kind = str(record.get("item_kind") or inferred_kind).strip().lower()
                item = ItemMaster(
                    sku=str(record.get("sku") or "").strip(),
                    manufacturer_sku=str(record.get("manufacturer_sku") or "").strip() or None,
                    description=str(record.get("description") or "").strip(),
                    manufacturer_dimensions=str(record.get("manufacturer_dimensions") or "").strip() or None,
                    flatpack_length=_parse_number(record.get("flatpack_length") or record.get("length")),
                    flatpack_width=_parse_number(record.get("flatpack_width") or record.get("width")),
                    flatpack_height=_parse_number(record.get("flatpack_height") or record.get("height")),
                    assembled_length=_parse_number(record.get("assembled_length")),
                    assembled_width=_parse_number(record.get("assembled_width")),
                    assembled_height=_parse_number(record.get("assembled_height")),
                    weight=_parse_number(record.get("weight")),
                    category=category,
                    item_kind=item_kind,
                    stackable=_parse_bool(record.get("stackable"), True),
                    fragile=_parse_bool(record.get("fragile"), False),
                    must_stay_upright=upright,
                    no_upside_down=_parse_bool(record.get("no_upside_down"), True),
                    allowed_rotations=_parse_rotations(record.get("allowed_rotations"), upright=upright),
                    max_stack_layers=(
                        int(value) if (value := _parse_number(record.get("max_stack_layers"))) is not None else None
                    ),
                    max_weight_on_top=_parse_number(record.get("max_weight_on_top"), 1000) or 0,
                    load_bearing_strength=_parse_number(record.get("load_bearing_strength"), 1000) or 0,
                    preferred_pallet_type=(str(record.get("preferred_pallet_type") or "either").strip().lower()),
                    freight_class=str(record.get("freight_class") or "").strip() or None,
                    nmfc=str(record.get("nmfc") or "").strip() or None,
                    color=str(record.get("color") or "").strip() or None,
                    rules_reviewed=_parse_bool(record.get("rules_reviewed"), False),
                    packing_role=str(record.get("packing_role") or "standard").strip().lower(),
                    bundle_max_units=int(
                        _parse_number(record.get("bundle_max_units"), 5 if item_kind == "accessory" else 1) or 1
                    ),
                    bundle_approved=_parse_bool(record.get("bundle_approved"), item_kind == "accessory"),
                    allow_edge_standing=_parse_bool(record.get("allow_edge_standing"), False),
                    overhang_approved=_parse_bool(record.get("overhang_approved"), False),
                    max_overhang_per_side=_parse_number(record.get("max_overhang_per_side"), 0) or 0,
                    active=_parse_bool(record.get("active"), True),
                )
                items.append(item)
                provided.append({RECORD_FIELDS.get(key, key) for key in record if key in ItemMaster.model_fields
                                 or key in RECORD_FIELDS})
            except Exception as exc:
                skipped.append(
                    {
                        "sheet": source_name,
                        "row": row_number,
                        "sku": str(record.get("sku") or "").strip(),
                        "reason": str(exc),
                    }
                )
    if not items:
        if skipped:
            preview = "; ".join(
                f"{row['sheet']} row {row['row']}: {row['reason']}" for row in skipped[:8]
            )
            raise ValueError(preview)
        raise ValueError("No valid item master rows were found")
    return ItemMasterImportResult(items=items, skipped=skipped, provided=provided)


def parse_cabinet_colors(filename: str, content: bytes) -> list[tuple[str, str]]:
    """(color name, item prefix) rows of a workbook's "Cabinet Colors" sheet (any sheet with those headers)."""

    if Path(filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    colors: list[tuple[str, str]] = []
    for _sheet, grid in _xlsx_grids(content):
        for index, row in enumerate(grid[:10]):
            keys = [_header_key(cell) for cell in row or []]
            if "itemprefix" in keys and ("cabinetcolor" in keys or "color" in keys):
                name_at = keys.index("cabinetcolor") if "cabinetcolor" in keys else keys.index("color")
                prefix_at = keys.index("itemprefix")
                for values in grid[index + 1:]:
                    values = list(values or [])
                    prefix = str(values[prefix_at] if prefix_at < len(values) and values[prefix_at] else "").strip()
                    name = str(values[name_at] if name_at < len(values) and values[name_at] else "").strip()
                    if prefix:
                        colors.append((name, prefix))
                break
    return colors


def parse_item_master(filename: str, content: bytes) -> list[ItemMaster]:
    """Compatibility wrapper for callers that only need valid item rows."""

    return parse_item_master_with_report(filename, content).items
