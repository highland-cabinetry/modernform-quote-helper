from io import BytesIO

from app.models import PackingSettings
from app.packing import StackPlannerPackingEngine
from app.services.reports import (
    packing_result_csv,
    packing_result_json,
    packing_result_pdf,
    packing_result_xlsx,
)
from openpyxl import load_workbook


def test_all_report_formats_include_warehouse_metadata(combined_factory):
    result = StackPlannerPackingEngine().pack(
        "Report test",
        [combined_factory(quantity=2)],
        PackingSettings(),
    )

    workbook = load_workbook(BytesIO(packing_result_xlsx(result)), data_only=True)
    assert "Pallets" in workbook.sheetnames
    assert "Layer Metrics" in workbook.sheetnames
    assert "Consolidation" in workbook.sheetnames
    assert "3D Placements" in workbook.sheetnames
    pallet_headers = [cell.value for cell in workbook["Pallets"][1]]
    assert "Largest cabinet-layer void %" in pallet_headers
    assert "Top support coverage %" in pallet_headers
    placement_headers = [cell.value for cell in workbook["3D Placements"][1]]
    assert "Bundle ID" in placement_headers
    assert "Support contacts" in placement_headers
    assert "Dunnage required" in placement_headers

    csv_text = packing_result_csv(result).decode("utf-8-sig")
    assert "bundle_id" in csv_text
    assert "requires_dunnage" in csv_text
    assert b'"layer_coverage_percent"' in packing_result_json(result)
    assert b'"largest_void_percent"' in packing_result_json(result)
    assert b'"top_support_coverage_percent"' in packing_result_json(result)
    assert b'"consolidation_attempts"' in packing_result_json(result)
    pdf = packing_result_pdf(result)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1_000
