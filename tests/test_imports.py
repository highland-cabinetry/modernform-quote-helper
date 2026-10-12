from io import BytesIO
from pathlib import Path

import pytest
from app.services.imports import (
    PdfPageData,
    _pdf_documents,
    parse_item_master_with_report,
    parse_packing_slip,
    parse_packing_slip_with_report,
)
from openpyxl import Workbook


def test_csv_packing_slip_import():
    content = b"Order,SKU,Description,Qty\nSO-1,WHS-B30,Base cabinet,2\n"
    lines = parse_packing_slip("slip.csv", content)
    assert lines[0].original_sku == "WHS-B30"
    assert lines[0].quantity == 2
    assert lines[0].source_ref == "slip.csv"


def test_xlsx_packing_slip_import():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Order 100"
    sheet.append(["Item Number", "Item Description", "Quantity"])
    sheet.append(["ESS-W3030", "Wall cabinet", 3])
    stream = BytesIO()
    workbook.save(stream)
    lines = parse_packing_slip("order.xlsx", stream.getvalue())
    assert lines[0].original_sku == "ESS-W3030"
    assert lines[0].quantity == 3
    assert "Order 100" in lines[0].source_ref


def test_shaker_item_master_headers_import_dual_dimensions_and_report_invalid_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Item References"
    sheet.append(
        [
            "In House Item SKU",
            "Manufacturer Item SKU",
            "Manufacturer Dimensions",
            "Unit Weight (lb)",
            "Flatpack L",
            "Flatpack W",
            "Flatpack H",
            "Assembled L",
            "Assembled W",
            "Assembled H",
            "Category",
            "Description",
        ]
    )
    sheet.append(
        [
            "B09", "M-B09", "9 x 24 x 34.5", 44.8419708, 36, 24.8, 3.8888889, 36, 10, 25,
            "Base cabinet", '9" Base Cabinet',
        ]
    )
    sheet.append(["BBC36", "", "", "", "", "", "", "", "", "", "Base cabinet", ""])
    stream = BytesIO()
    workbook.save(stream)

    result = parse_item_master_with_report("shaker.xlsx", stream.getvalue())

    assert len(result.items) == 1
    assert len(result.skipped) == 1
    item = result.items[0]
    assert item.sku == "B09"
    assert item.manufacturer_sku == "M-B09"
    assert item.flatpack_length == 36
    assert item.flatpack_width == 24.8
    assert item.flatpack_height == 3.8888889
    assert item.assembled_length == 36
    assert item.assembled_width == 10
    assert item.assembled_height == 25
    assert result.skipped[0]["sku"] == "BBC36"


def test_packing_state_column_defaults_to_flatpack_and_accepts_assembled():
    content = b"SKU,Qty,Packing State\nB09,1,assembled\nB12,1,\n"
    lines = parse_packing_slip("states.csv", content)
    assert lines[0].packing_state.value == "assembled"
    assert lines[1].packing_state.value == "flatpack"


def test_back_order_placeholders_preserve_authored_sku_and_description():
    content = (
        b"SKU,Qty,Description\n"
        b"BACK-ORDER,1,CSB42 corner sink base\n"
        b"BACK ORDER (Shaker),2,Trash Kit\n"
    )

    lines = parse_packing_slip("status.csv", content)

    assert [(line.original_sku, line.description) for line in lines] == [
        ("BACK-ORDER", "CSB42 corner sink base"),
        ("BACK ORDER (Shaker)", "Trash Kit"),
    ]


def test_packing_slip_report_ignores_non_items_and_keeps_unmatched_physical_rows():
    content = (
        b"Item,Quantity,Description\n"
        b"SBS,,****** Signature Blue Shaker *****\n"
        b"SBS-B09,2,Base cabinet\n"
        b"TRASH KIT,1,Trash kit\n"
        b"Discount,,-10 percent\n"
        b"Freight Handling ...,1,Pallet and labor\n"
        b"Freight(Shaker),,Freight terms\n"
    )

    result = parse_packing_slip_with_report("order.csv", content)

    assert [line.original_sku for line in result.lines] == ["SBS-B09", "TRASH KIT"]
    assert [line.quantity for line in result.lines] == [2, 1]
    assert [row["original_sku"] for row in result.ignored] == [
        "SBS", "Discount", "Freight Handling ...", "Freight(Shaker)",
    ]
    assert result.ignored[0]["reason_code"] == "SECTION_HEADING"
    assert all(row["source_ref"] == "order.csv" for row in result.ignored)


def test_uploaded_quantities_must_be_explicit_positive_whole_numbers():
    content = (
        b"SKU,Qty,Description\n"
        b"MISSING,,Missing\n"
        b"ZERO,0,Zero\n"
        b"NEGATIVE,-1,Negative\n"
        b"FRACTIONAL,1.5,Fractional\n"
        b"MALFORMED,each,Malformed\n"
        b"VALID,3,Valid\n"
    )

    result = parse_packing_slip_with_report("strict.csv", content)

    assert [(line.original_sku, line.quantity) for line in result.lines] == [("VALID", 3)]
    assert len(result.ignored) == 5
    assert result.ignored[0]["reason_code"] == "MISSING_QUANTITY"
    assert {row["reason_code"] for row in result.ignored[1:]} == {"INVALID_QUANTITY"}


def test_packing_slip_product_id_prefers_ship_quantity_then_ordered_quantity():
    content = (
        b"Qty,Product ID,Description,Ship Qty\n"
        b"7,WHS-B18,Partially shipped base,2\n"
        b"3,WHS-B24,Fallback ordered quantity,\n"
    )

    result = parse_packing_slip_with_report("packing-slip.csv", content)

    assert [(line.original_sku, line.quantity) for line in result.lines] == [
        ("WHS-B18", 2),
        ("WHS-B24", 3),
    ]


@pytest.mark.parametrize(
    ("path", "name", "line_count", "unit_count", "packing_state", "postal_code", "residential", "liftgate"),
    [
        (r"D:\Downloads\SO76325.pdf", "SO76325", 1, 1, "flatpack", "59714", False, False),
        (r"D:\Downloads\66557-I.pdf", "INV66557", 8, 11, "assembled", "87501", True, False),
        (r"D:\Downloads\66571 - p.pdf", "INV66571", 26, 41, "assembled", "79763", True, True),
    ],
)
def test_external_sales_order_invoice_and_packing_slip_regressions(
    path, name, line_count, unit_count, packing_state, postal_code, residential, liftgate
):
    source = Path(path)
    if not source.exists():
        pytest.skip("External customer regression PDF is intentionally not stored in the repository")
    result = parse_packing_slip_with_report(source.name, source.read_bytes())
    suggestion = result.shipping_suggestions[0]
    assert result.suggested_shipment_name == name
    assert len(result.lines) == line_count
    assert sum(line.quantity for line in result.lines) == unit_count
    assert {line.packing_state.value for line in result.lines} == {packing_state}
    assert suggestion["destination"]["postal_code"] == postal_code
    assert suggestion["accessorials"]["residential"] is residential
    assert suggestion["accessorials"]["liftgate"] is liftgate


def test_combined_pdf_metadata_keeps_document_context_and_blocks_conflicting_zip_autofill():
    pages = [
        PdfPageData(
            page_number=1,
            text="Highland Cabinetry Invoice\nSO75567\nRESIDENTIAL NO LIFTGATE",
            tables=[
                [["Date", "Invoice #"], ["8/17/2026", "66557"]],
                [["Ship To"], ["Customer\nSanta Fe, NM 87501"]],
                [["Ship Via", "Assembly"], ["Freight", "YES"]],
            ],
        ),
        PdfPageData(
            page_number=2,
            text="Packing Slip\nSO75699\nassemble freight - residential - add lift gate",
            tables=[
                [["Date", "Invoice #"], ["8/17/2026", "66571"]],
                [["Ship To"], ["Customer\nOdessa, TX 79763"]],
                [["Ship Via", "Assembly"], ["Freight", "YES"]],
            ],
        ),
    ]

    documents, suggestion, suggested_name = _pdf_documents("combined.pdf", pages)
    assert suggested_name == "INV66557 + INV66571"
    assert [document["primary_reference"] for document in documents] == ["INV66557", "INV66571"]
    assert suggestion is not None
    assert suggestion["destination"] == {}
    assert suggestion["destination_zip_candidates"] == ["87501", "79763"]
    assert any("different delivery ZIP" in warning for warning in suggestion["warnings"])
