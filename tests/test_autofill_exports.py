import io
import json

from app.models import BrokerRate, PackingSettings
from app.packing import StackPlannerPackingEngine
from app.services.reports import packing_result_json, packing_result_xlsx
from openpyxl import load_workbook


def test_old_rates_remain_in_model_but_are_not_exported(combined_factory):
    result = StackPlannerPackingEngine().pack("TEST", [combined_factory()], PackingSettings())
    result.broker_rates = [BrokerRate(broker="dsv", carrier="LEGACY-ONLY", price=123.45)]
    assert "broker_rates" not in json.loads(packing_result_json(result))
    workbook = load_workbook(io.BytesIO(packing_result_xlsx(result)))
    assert "Broker Rates" not in workbook.sheetnames
    assert result.broker_rates[0].carrier == "LEGACY-ONLY"
