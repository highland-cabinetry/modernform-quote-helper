"""Modernform edition: ModernForm invoices and packing slips → matched lines → pallet plan → broker packet.

The two PDFs in fixtures/modernform are a real ModernForm invoice (i4210_1.pdf) and the packing slip
of the same order (4210-P_Rev_1.pdf). The packing slip shortens long item codes ('NS-PC309024...').
"""

from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path
from urllib.parse import urlsplit

import pytest
from app import main
from app.database import SEED_MARKER, Repository
from app.edition import EDITIONS, seed_path
from app.models import MatchStatus, PackingSettings, PackingSlipLine
from app.services.adapters import ImportContext, choose_adapter
from app.services.sku_matching import match_lines
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures" / "modernform"
INVOICE = FIXTURES / "i4210_1.pdf"
PACKING_SLIP = FIXTURES / "4210-P_Rev_1.pdf"

# What both documents of order 4210 contain: printed code on the invoice → catalog SKU, quantity.
EXPECTED = {
    "NS-3DB30": ("3DB30", 1),
    "NS-CSB42": ("CSB42", 3),
    "NS-PC309024BOX": ("PC309024BOX", 8),
    "NS-PC309024FFD": ("PC309024FFD", 8),
    "NS-PC309624BOX": ("PC309624BOX", 6),
    "NS-PC309624FFD": ("PC309624FFD", 6),
    "NS-PNL34.5x96x1/4": ("PNL 34.5x96x1/4", 3),
    "NS-OCM8": ("OCM8", 40),
    "NS-SDM8": ("SDM8", 30),
    "NS-TF1.5x96": ("TF1.5x96", 70),
    "NS-TF6x96": ("TF6x96", 40),
    "NS-VSD3621DR": ("VSD3621DR", 5),
    "NS-WF336": ("WF336", 50),
    "Trash Kit Shaker": ("TRASH KIT", 20),
    "WS-Sample Door": ("Sample door", 50),
    "NO-Sample Door": ("Sample door", 50),
    "SDW-Sample Door": ("Sample door", 100),
}
TOTAL_UNITS = sum(quantity for _, quantity in EXPECTED.values())


@pytest.fixture
def modernform(monkeypatch):
    monkeypatch.setenv("QUOTE_HELPER_EDITION", "modernform")
    monkeypatch.delenv("QUOTE_HELPER_SEED_CATALOG", raising=False)


@pytest.fixture
def repo(modernform, tmp_path) -> Repository:
    repository = Repository(tmp_path / "modernform.sqlite3", tmp_path / "backups")
    repository.initialize()
    return repository


@pytest.fixture
def client(repo, monkeypatch):
    monkeypatch.setattr(main, "repository", repo)
    with TestClient(main.app) as test_client:
        yield test_client


def upload(client: TestClient, *paths: Path) -> dict:
    response = client.post("/api/imports", files=[("files", (path.name, path.read_bytes(), "application/pdf"))
                                                  for path in paths])
    assert response.status_code == 200, response.text
    return response.json()


def parsed_lines(repo: Repository, path: Path) -> list[PackingSlipLine]:
    company = repo.default_company()
    context = ImportContext(filename=path.name, content=path.read_bytes(), config=company.config)
    adapter, score = choose_adapter(context, company.config.adapters)
    assert adapter.name == "modernform_pdf" and score > 0.9
    outcome = adapter.parse(context)
    assert outcome.status == "ok" and len(outcome.orders) == 1
    return [PackingSlipLine(original_sku=line.original_sku, description=line.description, quantity=line.quantity,
                            packing_state=line.packing_state, source_ref=line.source_ref)
            for line in outcome.orders[0].lines]


# --------------------------------------------------------------------------- edition and starter data


def test_a_new_database_belongs_to_modernform_with_the_starter_catalog(repo):
    company = repo.default_company()
    assert (company.code, company.name) == ("MCC", "ModernForm Cabinets & Closet")
    assert company.config.rule_set == "modernform"
    assert company.config.adapters[0] == "modernform_pdf"
    assert company.config.origin.postal_code == "75006"
    seed = json.loads(seed_path(EDITIONS["modernform"]).read_text(encoding="utf-8"))
    assert repo.item_count() == len(seed["items"]) > 300
    settings = repo.get_settings()
    assert settings.quote_origin.postal_code == "75006" and settings.quote_origin.state == "TX"
    assert settings.standard_pallet.max_total_weight == 1200 and settings.long_pallet.max_total_weight == 1600
    assert not any("Highland" in item.description for item in repo.list_items(active_only=False))


def test_the_starter_catalog_never_overwrites_or_duplicates(repo):
    item = next(item for item in repo.list_items() if item.sku == "3DB30")
    item.weight = 123.0
    repo.upsert_items([item])
    repo.delete_item(next(item.id for item in repo.list_items() if item.sku == "B09"))
    settings = repo.get_settings()
    settings.quote_origin.postal_code = "75001"
    repo.save_settings(settings)
    count = repo.item_count()

    repo.initialize()  # every start
    assert repo.apply_edition_seed()["applied"] is False
    assert repo.item_count() == count
    assert next(item for item in repo.list_items() if item.sku == "3DB30").weight == 123.0
    assert "B09" not in {item.sku for item in repo.list_items()}
    assert repo.get_settings().quote_origin.postal_code == "75001"


def test_a_newer_seed_version_adds_only_missing_products(repo, tmp_path, monkeypatch):
    seed = json.loads(seed_path(EDITIONS["modernform"]).read_text(encoding="utf-8"))
    seed["seed_version"] += 1
    changed = dict(seed["items"][0], weight=999.0)
    new_item = dict(seed["items"][0], sku="NEW-TEST-SKU", description="Added by a later release")
    seed["items"] = [changed, new_item]
    later = tmp_path / "later_seed.json"
    later.write_text(json.dumps(seed), encoding="utf-8")
    monkeypatch.setattr("app.edition.seed_path", lambda edition=None: later)
    before = repo.item_count()
    result = repo.apply_edition_seed()
    assert result["applied"] and result["items_added"] == 1
    assert repo.item_count() == before + 1
    assert next(item for item in repo.list_items() if item.sku == changed["sku"]).weight != 999.0
    with repo.connect() as connection:
        marker = json.loads(connection.execute("SELECT value_json FROM settings WHERE key=?",
                                               (SEED_MARKER,)).fetchone()[0])
    assert marker["seed_version"] == seed["seed_version"]


def test_the_starter_catalog_can_be_turned_off(modernform, tmp_path, monkeypatch):
    monkeypatch.setenv("QUOTE_HELPER_SEED_CATALOG", "0")
    repository = Repository(tmp_path / "empty.sqlite3", tmp_path / "backups")
    repository.initialize()
    assert repository.item_count() == 0


# --------------------------------------------------------------------------- documents and matching


@pytest.mark.parametrize("path", [INVOICE, PACKING_SLIP], ids=["invoice", "packing-slip"])
def test_every_line_of_both_documents_matches_the_right_product(repo, path):
    lines = parsed_lines(repo, path)
    assert len(lines) == len(EXPECTED) and sum(line.quantity for line in lines) == TOTAL_UNITS
    settings = repo.get_settings()
    matches = match_lines(lines, repo.catalog_items("default"), ignored_prefixes=settings.sku_prefixes_to_ignore,
                          fuzzy_threshold=settings.fuzzy_match_threshold,
                          confirmed_aliases=repo.catalog_alias_targets("default"), rule_set="modernform")
    found: dict[str, tuple[str, int]] = {}
    for match in matches:
        assert match.match_status in {MatchStatus.EXACT, MatchStatus.CONFIRMED_ALIAS}, match.review_message
        assert match.matched_item is not None
        found.setdefault(match.matched_item.sku, (match.matched_item.sku, 0))
        found[match.matched_item.sku] = (match.matched_item.sku, found[match.matched_item.sku][1] + match.line.quantity)
    expected: dict[str, int] = {}
    for sku, quantity in EXPECTED.values():
        expected[sku] = expected.get(sku, 0) + quantity
    assert {sku: quantity for sku, (_, quantity) in found.items()} == expected
    if path == PACKING_SLIP:
        shortened = [match for match in matches if match.line.original_sku.endswith("...")]
        assert len(shortened) == 10   # 9 shortened codes + 'Trash Kit Sha...'
        assert all(match.match_status == MatchStatus.CONFIRMED_ALIAS and "verify" in match.review_message
                   for match in shortened)


def test_header_fields_are_read(repo):
    company = repo.default_company()
    for path, document_type in ((INVOICE, "invoice"), (PACKING_SLIP, "packing_slip")):
        context = ImportContext(filename=path.name, content=path.read_bytes(), config=company.config)
        adapter, _ = choose_adapter(context, company.config.adapters)
        [order] = adapter.parse(context).orders
        assert order.order_number == "INV4210" and order.document_type == document_type
        assert order.ship_to["postal_code"] == "80011" and order.ship_to["state"] == "CO"
        assert order.customer_name == "Highland Cabinetry Colorado 01"
        assert order.sales_rep == "IT" and order.order_date == "8/26/2026"
        assert order.packing_state.value == "flatpack"
        # Delivery/return notices printed in the item table are not items and do not flag the order.
        assert all(row["reason_code"] == "NON_PHYSICAL_ROW" for row in order.ignored)
    assert "SO13157" in adapter.parse(ImportContext(filename=INVOICE.name, content=INVOICE.read_bytes(),
                                                    config=company.config)).orders[0].references


def test_ambiguous_shortened_codes_are_held_for_review(repo):
    line = PackingSlipLine(original_sku="NS-PC309024...", description="Pantry part", quantity=1, source_ref="t")
    settings = repo.get_settings()
    [match] = match_lines([line], repo.catalog_items("default"), ignored_prefixes=settings.sku_prefixes_to_ignore,
                          fuzzy_threshold=settings.fuzzy_match_threshold, confirmed_aliases={},
                          rule_set="modernform")
    assert match.match_status == MatchStatus.UNMATCHED and match.matched_item is None
    assert "PC309024BOX" in match.review_message and "PC309024FFD" in match.review_message


def test_assemble_column_marks_assembled_lines(repo):
    from app.services.adapters.table import header_aliases, rows_to_lines

    rows = [["Item", "Description", "Assemble", "Quantity"],
            ["NS-3DB30", "30\" 3 Drawer Base", "Yes", "1"],
            ["NS-CSB42", "42\" Cabinet Sink Base", "", "3"]]
    pairs, _, found = rows_to_lines(rows, "slip", header_aliases(repo.default_company().config), None)
    assert found
    assert [line.packing_state.value if line.packing_state else None for _, line in pairs] == ["assembled", None]


# --------------------------------------------------------------------------- the whole workflow


def test_import_both_documents_pack_and_prepare_every_broker(client, repo):
    # Packing slip first: the order is created with shortened codes ...
    first = upload(client, PACKING_SLIP)
    assert first["summary"]["created"] == 1 and first["summary"]["conflicts"] == 0
    [order] = first["orders"]
    assert order["order_number"] == "INV4210" and order["status"] == "ready", order["issues"]
    assert sum(line["quantity"] for line in order["lines"]) == TOTAL_UNITS

    # ... then the invoice of the same order: same goods, so no conflict; the full codes are kept.
    second = upload(client, INVOICE)
    assert second["summary"]["conflicts"] == 0 and second["summary"]["created"] == 0
    stored = client.get(f"/api/orders/{order['id']}").json()["order"]
    assert sorted(line["original_sku"] for line in stored["lines"]) == sorted(EXPECTED)
    assert "SO13157" in stored["references"]

    selection = client.post("/api/orders/packing-lines", json={"order_ids": [order["id"]]}).json()
    assert selection["settings"]["quote_origin"]["postal_code"] == "75006"
    request = {"shipment_name": selection["suggested_name"], "lines": selection["lines"],
               "settings": selection["settings"], "company_id": selection["company_id"],
               "order_ids": [order["id"]]}
    response = client.post("/api/shipments/pack", json=request)
    assert response.status_code == 200, response.text
    plan = response.json()
    assert plan["pallets"], plan.get("warnings")
    assert not plan["unpacked"], plan["unpacked"][:3]
    errors = [issue for issue in plan["warnings"] if issue["severity"] == "error"]
    assert not errors, errors
    # 14 pantry BOX/FFD pairs used to take one long pallet each (14 + 3); now the order needs six.
    assert len(plan["pallets"]) <= 6, [(p["pallet_type"], p["freight_height"]) for p in plan["pallets"]]
    limits = selection["settings"]
    for pallet in plan["pallets"]:
        limit = limits["long_pallet" if pallet["pallet_type"] == "long" else "standard_pallet"]["max_total_weight"]
        assert pallet["total_weight"] <= limit + 1e-6

    shipping = {"confirmed": True, "customer_references": ["INV4210"],
                "origin": {"postal_code": "75006", "city": "Carrollton", "state": "TX"},
                "destination": {"postal_code": "80011", "city": "Aurora", "state": "CO"},
                "accessorials": {"residential": False, "liftgate": False, "appointment_delivery": True}}
    saved = client.put(f"/api/shipments/{plan['id']}/shipping-info", json=shipping)
    assert saved.status_code == 200, saved.text
    preview = client.post("/api/quote-handoffs/preview", json={"shipment_id": plan["id"]}).json()
    assert preview["ready"], preview
    launches = client.post("/api/quote-handoffs/prepare-all", json={"shipment_id": plan["id"]})
    assert launches.status_code == 200, launches.text
    handoffs = launches.json()["handoffs"]
    try:
        check_packet(client, plan, handoffs, launches.json()["skipped"])
    finally:   # handoffs live in the app process; leave none active for other tests
        for handoff in handoffs:
            client.post(f"/api/quote-handoffs/{handoff['request_id']}/cancel")


def check_packet(client, plan, handoffs, skipped):
    # ModernForm quotes with TQL, DSV and Old Dominion only (EchoShip and MyUnishippers are off). Every
    # enabled broker opens, except one that cannot take this shipment (Old Dominion: pallets over 96 in
    # long), which is reported instead of stopping the others.
    opened = {handoff["broker"] for handoff in handoffs}
    assert opened | {entry["broker"] for entry in skipped} == {"tql", "dsv", "odfl"}
    assert {"tql", "dsv"} <= {handoff["broker"] for handoff in handoffs}
    for broker in ("echoship", "myunishippers"):
        refused = client.post(f"/api/quote-handoffs/{broker}/prepare", json={"shipment_id": plan["id"]})
        assert refused.status_code == 422 and "turned off" in refused.text
    assert all(entry["broker"] == "odfl" and "96" in entry["reason"] for entry in skipped)

    # The Modernform Bridge (its own extension ID) can read the packet; another extension cannot.
    fragment = urlsplit(handoffs[0]["url"]).fragment.split("build2quote=", 1)[1]
    launch = json.loads(base64.urlsafe_b64decode(fragment + "=" * (-len(fragment) % 4)))
    path = f"/api/quote-handoffs/{launch['request_id']}/packet"
    packet = client.get(path, params={"token": launch["token"]}, headers={"Origin": main.BUILD2QUOTE_EXTENSION_ORIGIN})
    assert packet.status_code == 200, packet.text
    body = packet.json()["packet"]
    assert body["origin"]["postal_code"] == "75006" and body["destination"]["postal_code"] == "80011"
    assert len(body["pallets"]) == len(plan["pallets"])
    other = client.get(path, params={"token": launch["token"]},
                       headers={"Origin": "chrome-extension://celfkhbiefaheonffgddnocaljcbekik"})
    assert other.status_code == 403


# --------------------------------------------------------------------------- separate identity


def test_bridge_extension_id_matches_its_manifest_key():
    manifest = json.loads((ROOT / "build2quote_bridge/manifest.json").read_text(encoding="utf-8"))
    digest = hashlib.sha256(base64.b64decode(manifest["key"])).hexdigest()[:32]
    extension_id = "".join(chr(ord("a") + int(character, 16)) for character in digest)
    assert extension_id == main.BUILD2QUOTE_EXTENSION_ID
    frontend = (ROOT / "frontend/src/build2quote.js").read_text(encoding="utf-8")
    assert f'"{extension_id}"' in frontend
    assert extension_id != "celfkhbiefaheonffgddnocaljcbekik"  # Highland's Bridge
    assert manifest["name"].startswith("Modernform")


def test_installed_app_is_separate_from_highland_quote_helper():
    import sys

    sys.path.insert(0, str(ROOT / "desktop"))
    import desktop_common as common
    import desktop_update as updates

    assert common.APP_ID == updates.PRODUCT == "ModernformQuoteHelper"
    assert not set(common.PORT_RANGE) & set(range(8000, 8010))
    script = (ROOT / "desktop/installer/QuoteHelper.iss").read_text(encoding="utf-8")
    assert "894D1B90-1F45-4388-BEF9-587E6CE0107C" not in script  # Highland's AppId
    assert f'"port": {common.DEFAULT_PORT}' in script
    release = json.loads((ROOT / "release.json").read_text(encoding="utf-8"))
    assert release["edition"] == "modernform"


def test_settings_default_prefixes_still_validate():
    PackingSettings.model_validate(json.loads(seed_path(EDITIONS["modernform"]).read_text("utf-8"))["settings"])


def test_pantry_pairs_stack_several_to_a_long_pallet(repo):
    """8 PC309024 and 6 PC309624 BOX/FFD pairs: boxes stacked, their FFDs on top, every pair kept together."""
    from app.models import PackingRequest
    from app.services.shipments import build_shipment

    lines = [line for line in parsed_lines(repo, INVOICE) if "PC30" in line.original_sku]
    result = build_shipment(PackingRequest(shipment_name="pantries", lines=lines, settings=repo.get_settings(),
                                           save_result=False), repo)
    assert not result.unpacked
    assert not [issue for issue in result.warnings if issue.severity == "error"], result.warnings
    assert len(result.pallets) == 4 and all(pallet.pallet_type == "long" for pallet in result.pallets)
    for pallet in result.pallets:
        families = {item.sku.removesuffix("BOX").removesuffix("FFD") for item in pallet.items}
        boxes = [item for item in pallet.items if item.sku.endswith("BOX")]
        ffds = [item for item in pallet.items if item.sku.endswith("FFD")]
        assert len(families) == 1 and len(boxes) == len(ffds)
        assert max(box.z for box in boxes) < min(ffd.z for ffd in ffds)   # every FFD above every BOX


# --------------------------------------------------------------------------- simple catalog workbook


def test_catalog_workbook_round_trip_keeps_handling_rules_and_colors(client, repo):
    import io

    from openpyxl import load_workbook

    exported = client.get("/api/catalog/default/export/xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook.sheetnames == ["Cabinet Colors", "Item References"]
    items = workbook["Item References"]
    assert [cell.value for cell in items[1]] == [
        "In House Item SKU", "Manufacturer Item SKU", "Manufacturer Dimensions", "Unit Weight (lb)", "Flatpack L",
        "Flatpack W", "Flatpack H", "Assembled L", "Assembled W", "Assembled H", "Category", "Description"]
    assert [[cell.value for cell in row] for row in workbook["Cabinet Colors"].iter_rows()] == [
        ["Cabinet Color", "Item Prefix"], ["Natural Shaker", "NS-"]]

    def preview(content: bytes) -> dict:
        response = client.post("/api/catalog/default/import/preview",
                               files={"file": ("catalog.xlsx", content, "application/octet-stream")})
        assert response.status_code == 200, response.text
        return response.json()

    unchanged = preview(exported.content)
    assert unchanged["counts"]["changed"] == unchanged["counts"]["new"] == 0 and not unchanged["can_apply"]

    for row in items.iter_rows(min_row=2):
        if row[0].value == "PC309024FFD":
            row[3].value = 42.5
    items.append(["GS-TEST", None, None, 50, 30, 20, 4, 0, 0, 0, "Base cabinet", "Test base"])
    workbook["Cabinet Colors"].append(["Gray Slate", "gs"])
    edited = io.BytesIO()
    workbook.save(edited)
    result = preview(edited.getvalue())
    assert result["counts"]["changed"] == 1 and result["counts"]["new"] == 1
    assert {color["prefix"] for color in result["colors"]} == {"NS-", "GS-"}
    applied = client.post("/api/catalog/default/import/apply", json={"preview_id": result["preview_id"]})
    assert applied.status_code == 200, applied.text

    ffd = next(item for item in repo.list_items() if item.sku == "PC309024FFD")
    assert ffd.weight == 42.5 and ffd.rules_reviewed and ffd.max_weight_on_top == 100 and ffd.fragile
    new = next(item for item in repo.list_items() if item.sku == "GS-TEST")
    assert new.assembled_length is None and new.flatpack_length == 30
    settings = repo.get_settings()
    assert [color.prefix for color in settings.cabinet_colors] == ["NS-", "GS-"]
    assert "GS" in settings.sku_prefixes_to_ignore


def test_replacing_the_catalog_keeps_settings_and_never_brings_back_the_starter_catalog(client, repo):
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    colors = workbook.active
    colors.title = "Cabinet Colors"
    colors.append(["Cabinet Color", "Item Prefix"])
    colors.append(["Walnut Slab", "WS-"])
    items = workbook.create_sheet("Item References")
    items.append(["In House Item SKU", "Unit Weight (lb)", "Flatpack L", "Flatpack W", "Flatpack H", "Category",
                  "Description"])
    items.append(["WB30", 70, 36, 31, 5, "Base cabinet", "30 base"])
    content = io.BytesIO()
    workbook.save(content)
    response = client.post("/api/item-master/replace", data={"confirmation": "REPLACE"},
                           files={"file": ("ours.xlsx", content.getvalue(), "application/octet-stream")})
    assert response.status_code == 200, response.text
    assert [item.sku for item in repo.list_items()] == ["WB30"]
    settings = repo.get_settings()
    assert settings.quote_origin.postal_code == "75006"
    assert [color.prefix for color in settings.cabinet_colors] == ["WS-"]
    repo.initialize()   # next start
    assert [item.sku for item in repo.list_items()] == ["WB30"]
