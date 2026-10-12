import json

from app import main
from app.database import Repository
from app.models import ContainedItem, PackingSettings
from app.packing import StackPlannerPackingEngine
from app.services.broker_handoff import _required_missing
from fastapi.testclient import TestClient


def test_missing_or_duplicated_units_block_freight(combined_factory):
    result = StackPlannerPackingEngine().pack("account", [combined_factory(quantity=2)], PackingSettings())
    assert result.freight_ready
    result.pallets[0].items.pop()
    assert result.accounting_errors and not result.freight_ready
    assert any("freight-ready" in value for value in _required_missing(None, result))


def test_contained_host_weight_and_capacity_are_verified(combined_factory):
    result = StackPlannerPackingEngine().pack("account", [combined_factory(sku="B18")], PackingSettings())
    host = result.pallets[0].items[0]
    result.contained_items = [
        ContainedItem(
            unit_id="kit",
            host_unit_id=host.unit_id,
            sku="TRASH KIT",
            original_sku="TRASH KIT",
            weight=18,
            host_base_weight=host.weight,
        )
    ]
    result.summary.total_items += 1
    result.summary.packed_items += 1
    assert not result.freight_ready  # flatpack host and kit weight absent


def test_saved_incomplete_ready_claim_cannot_export_or_quote(tmp_path, monkeypatch, combined_factory):
    repo = Repository(tmp_path / "isolated.sqlite3")
    repo.initialize()
    monkeypatch.setattr(main, "repository", repo)
    result = StackPlannerPackingEngine().pack("account", [combined_factory(quantity=2)], PackingSettings())
    raw = result.model_dump(mode="json")
    raw["pallets"][0]["items"].pop()
    # Persist malformed legacy evidence, retaining its inaccurate ready claim.
    with repo.connect() as db:
        db.execute(
            "INSERT INTO shipments(id,shipment_name,request_json,result_json,created_at) VALUES (?,?,?,?,?)",
            (result.id, "account", "{}", json.dumps(raw), result.created_at.isoformat()),
        )
    with TestClient(main.app) as client:
        for fmt in ["pdf", "xlsx", "csv", "json"]:
            assert client.get(f"/api/shipments/{result.id}/export/{fmt}").status_code == 422
        response = client.post("/api/quote-handoffs/tql/prepare", json={"shipment_id": result.id})
        assert response.status_code in (400, 409, 422)


def test_assembled_kit_reconciles_source_counts_and_reports():
    import io

    from app.models import ItemMaster, PackingRequest, PackingSlipLine, PackingState
    from app.packing_worker import SnapshotRepository
    from app.services.reports import packing_result_csv, packing_result_pdf, packing_result_xlsx
    from app.services.shipments import build_shipment
    from openpyxl import load_workbook

    host = ItemMaster(
        id=18,
        sku="B18",
        assembled_length=36,
        assembled_width=19,
        assembled_height=25,
        weight=60,
        item_kind="cabinet",
        category="Base cabinet",
    )
    request = PackingRequest(
        shipment_name="contained",
        save_result=False,
        lines=[
            PackingSlipLine(original_sku="B18", packing_state=PackingState.ASSEMBLED, quantity=1, source_ref="TEST"),
            PackingSlipLine(
                original_sku="Back-Order(Shaker)",
                description="TRASH KIT",
                packing_state=PackingState.ASSEMBLED,
                quantity=1,
                source_ref="TEST",
            ),
        ],
    )
    repo = SnapshotRepository({"master": [host.model_dump(mode="json")], "aliases": {}})
    result = build_shipment(request, repo)
    assert result.freight_ready, result.accounting_errors
    assert result.summary.total_items == result.summary.packed_items == 2
    assert len(result.contained_items) == 1
    assert sum(len(p.items) for p in result.pallets) == 1
    assert result.summary.packed_item_weight == 78
    assert result.contained_items[0].weight == 18
    assert "inside B18" in packing_result_csv(result).decode("utf-8-sig")
    book = load_workbook(io.BytesIO(packing_result_xlsx(result)))
    assert book["Contained items"].cell(2, 4).value == "inside B18"
    assert packing_result_pdf(result).startswith(b"%PDF")
    prefixed = result.model_copy(deep=True)
    prefixed.pallets[0].items[0].sku = "WHS-B18"
    assert prefixed.freight_ready
    invalid_content = result.model_copy(deep=True)
    invalid_content.contained_items[0].sku = "WBC3042"
    assert not invalid_content.freight_ready
    duplicate = result.model_copy(deep=True)
    duplicate.contained_items.append(duplicate.contained_items[0])
    assert not duplicate.freight_ready


def test_failed_host_leaves_kit_explicitly_unpacked():
    from app.models import ItemMaster, PackingRequest, PackingSlipLine
    from app.packing_worker import SnapshotRepository
    from app.services.shipments import build_shipment

    host = ItemMaster(id=18, sku="B18", assembled_length=150, assembled_width=150, assembled_height=25, weight=60)
    request = PackingRequest(
        shipment_name="host blocked",
        save_result=False,
        lines=[
            PackingSlipLine(original_sku="B18", packing_state="assembled", source_ref="TEST"),
            PackingSlipLine(original_sku="TRASH KIT", packing_state="assembled", source_ref="TEST"),
        ],
    )
    result = build_shipment(request, SnapshotRepository({"master": [host.model_dump(mode="json")], "aliases": {}}))
    assert not result.freight_ready
    assert len(result.unpacked) == result.summary.unpacked_items == 2
    assert result.summary.total_items == 2
    assert result.summary.total_item_weight == 78
    assert not result.contained_items


def test_stale_summary_fields_block_freight(combined_factory):
    result = StackPlannerPackingEngine().pack("Summary", [combined_factory()], PackingSettings())
    for field in ["total_item_weight", "total_freight_weight", "total_pallets", "pallet_types"]:
        changed = result.model_copy(deep=True)
        setattr(changed.summary, field, {} if field == "pallet_types" else 999)
        assert not changed.freight_ready, field
